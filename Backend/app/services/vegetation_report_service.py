"""
Vegetation Indices Report generator - permanent fixed version.
- Correctly finds PNG files on Windows
- Handles odd number of images cleanly (last image always visible)
- Clean 2-column layout
"""

from app.models.ndvi_history import NdviHistory
from sqlalchemy.orm import Session
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx import Document
import matplotlib.pyplot as plt
import io
import os
from datetime import date, datetime, timedelta
from urllib.parse import urlparse
from pathlib import Path

import matplotlib
matplotlib.use("Agg")


INDEX_ORDER = [
    ("ndvi", "NDVI", "#2D6A4F"),
    ("ndmi", "NDMI", "#1D4E89"),
    ("ndre", "NDRE", "#B45309"),
    ("evi", "EVI", "#40916C"),
    ("savi", "SAVI", "#95D5B2"),
    ("exg", "ExG", "#65A30D"),
    ("vari", "VARI", "#0D9488"),
    ("gli", "GLI", "#CA8A04"),
]

_LEFTOVER_MERGE_THRESHOLD = 5
_INNER_TOLERANCE = 2
_EDGE_TOLERANCE = 7


def _url_to_local_path(url: str) -> str | None:
    """Convert a PNG URL into a reliable absolute local path (Windows + Linux)."""
    if not url:
        return None

    path = urlparse(url).path
    marker = "/static/"
    idx = path.find(marker)
    if idx == -1:
        return None

    relative = path[idx + len(marker):]  # ndvi_images/xxx.png

    current_file = Path(__file__).resolve()
    project_root = current_file.parents[2]

    candidates = [
        project_root / "static" / relative,
        Path.cwd() / "static" / relative,
        Path("static") / relative,
        Path(relative),
    ]

    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return str(candidate)

    return None


def _build_weekly_trend(rows: list[NdviHistory], index_key: str, start: date, end: date) -> list[dict]:
    mean_attr = f"{index_key}_mean"

    total_days = (end - start).days + 1
    full_weeks = total_days // 7
    leftover = total_days - full_weeks * 7

    buckets: list[tuple[date, date]] = []
    if full_weeks == 0:
        buckets.append((start, end))
    else:
        cursor = start
        for i in range(full_weeks):
            is_last_full_week = i == full_weeks - 1
            bucket_end = cursor + timedelta(days=6)
            if is_last_full_week and 0 < leftover < _LEFTOVER_MERGE_THRESHOLD:
                bucket_end = end
            buckets.append((cursor, bucket_end))
            cursor = bucket_end + timedelta(days=1)
        if leftover >= _LEFTOVER_MERGE_THRESHOLD:
            buckets.append((cursor, end))

    used: set[int] = set()

    def has_value(r: NdviHistory) -> bool:
        return getattr(r, mean_attr, None) is not None

    def summarize(matched: list[NdviHistory]) -> dict:
        values = [getattr(r, mean_attr) for r in matched]
        return {
            "mean": round(sum(values) / len(values), 4),
            "max": round(max(values), 4),
            "min": round(min(values), 4),
            "count": len(values),
        }

    exact_matches: list[list[int]] = []
    for b_start, b_end in buckets:
        matched = [
            i for i, r in enumerate(rows)
            if i not in used and has_value(r) and b_start <= r.satellite_image_date <= b_end
        ]
        for i in matched:
            used.add(i)
        exact_matches.append(matched)

    final_matches: list[list[int]] = []
    for idx, (b_start, b_end) in enumerate(buckets):
        if exact_matches[idx]:
            final_matches.append(exact_matches[idx])
            continue
        is_first = idx == 0
        is_last = idx == len(buckets) - 1
        before_tol = _EDGE_TOLERANCE if is_first else _INNER_TOLERANCE
        after_tol = _EDGE_TOLERANCE if is_last else _INNER_TOLERANCE
        search_start = b_start - timedelta(days=before_tol)
        search_end = b_end + timedelta(days=after_tol)
        matched = [
            i for i, r in enumerate(rows)
            if i not in used and has_value(r) and search_start <= r.satellite_image_date <= search_end
        ]
        for i in matched:
            used.add(i)
        final_matches.append(matched)

    output = []
    for (b_start, b_end), matched in zip(buckets, final_matches):
        matched_rows = [rows[i] for i in matched]
        if not matched_rows:
            output.append({"week_start": b_start, "week_end": b_end,
                          "mean": None, "count": 0, "rows": []})
            continue
        point = {"week_start": b_start,
                 "week_end": b_end, "rows": matched_rows}
        point.update(summarize(matched_rows))
        output.append(point)
    return output


def _build_drone_series(rows: list[NdviHistory], index_key: str) -> list[dict]:
    mean_attr = f"{index_key}_mean"
    with_value = [r for r in rows if getattr(r, mean_attr, None) is not None]
    seen_dates = set()
    deduped = []
    for r in sorted(with_value, key=lambda r: r.satellite_image_date):
        if r.satellite_image_date in seen_dates:
            continue
        seen_dates.add(r.satellite_image_date)
        deduped.append(r)

    return [
        {
            "week_start": r.satellite_image_date,
            "week_end": r.satellite_image_date,
            "mean": getattr(r, mean_attr),
            "max": getattr(r, mean_attr),
            "min": getattr(r, mean_attr),
            "count": 1,
            "rows": [r],
        }
        for r in deduped
    ]


def _render_trend_chart(points: list[dict], index_label: str, color: str) -> bytes | None:
    plotted = [p for p in points if p["mean"] is not None]
    if not plotted:
        return None

    fig, ax = plt.subplots(figsize=(6.4, 2.5), dpi=140)
    dates = [p["week_start"] for p in plotted]
    values = [p["mean"] for p in plotted]

    ax.plot(dates, values, color=color, linewidth=2.2,
            marker="o", markersize=5, markerfacecolor="white", markeredgewidth=1.4)
    ax.set_title(f"{index_label} trend", fontsize=12, fontweight="bold",
                 color="#1e2b23", loc="left", pad=6)
    ax.tick_params(axis="both", labelsize=8, colors="#6b7280")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#d1d5db")
    ax.spines["bottom"].set_color("#d1d5db")
    ax.grid(axis="y", linestyle="--", alpha=0.35)
    fig.autofmt_xdate(rotation=25)
    fig.tight_layout(pad=0.5)

    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def _add_index_section(document: Document, index_key: str, index_label: str, color: str, points: list[dict]) -> None:
    # Colored heading
    heading = document.add_heading(index_label, level=2)
    for run in heading.runs:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        run.font.size = Pt(14)

    # Trend chart
    chart_bytes = _render_trend_chart(points, index_label, color)
    if chart_bytes:
        document.add_picture(io.BytesIO(chart_bytes), width=Inches(6.2))
        document.add_paragraph()
    else:
        p = document.add_paragraph("No trend data available for this index.")
        p.runs[0].italic = True
        document.add_paragraph()

    # Collect valid images
    valid = []
    for point in points:
        if point["count"] == 0 or not point["rows"]:
            continue
        row = point["rows"][0]
        png_url = getattr(row, f"{index_key}_png_url", None)
        local_path = _url_to_local_path(png_url) if png_url else None

        if local_path and os.path.exists(local_path) and os.path.getsize(local_path) > 200:
            valid.append((point, local_path))

    if not valid:
        p = document.add_paragraph(
            "No heatmap images available for this index.")
        p.runs[0].italic = True
        document.add_paragraph()
        return

    # Add images – correctly handle odd numbers
    i = 0
    while i < len(valid):
        remaining = len(valid) - i

        if remaining == 1:
            # Last single image → put it alone (centered)
            point, path = valid[i]

            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            try:
                run.add_picture(path, width=Inches(3.2))
            except Exception:
                pass

            w_start = point["week_start"].strftime("%d-%m-%Y")
            w_end = point["week_end"].strftime("%d-%m-%Y")
            date_label = w_start if w_start == w_end else f"{w_start} – {w_end}"

            cap = document.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap.add_run(
                f"{date_label}\nMean {point['mean']:.3f}  ·  Max {point['max']:.3f}  ·  Min {point['min']:.3f}"
            )
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

            document.add_paragraph()
            i += 1
        else:
            # Two images side by side
            table = document.add_table(rows=1, cols=2)
            table.autofit = True

            for col in range(2):
                cell = table.cell(0, col)
                cell.paragraphs[0].clear()

                point, path = valid[i + col]

                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                try:
                    run.add_picture(path, width=Inches(2.95))
                except Exception:
                    continue

                w_start = point["week_start"].strftime("%d-%m-%Y")
                w_end = point["week_end"].strftime("%d-%m-%Y")
                date_label = w_start if w_start == w_end else f"{w_start} – {w_end}"

                cap = cell.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(
                    f"{date_label}\nMean {point['mean']:.3f}  ·  Max {point['max']:.3f}  ·  Min {point['min']:.3f}"
                )
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

            document.add_paragraph()
            i += 2


def build_vegetation_report_docx(db: Session, field_id, field_name: str) -> bytes:
    all_rows: list[NdviHistory] = (
        db.query(NdviHistory)
        .filter(NdviHistory.field_id == field_id)
        .order_by(NdviHistory.satellite_image_date.asc())
        .all()
    )
    satellite_rows = [r for r in all_rows if r.source_collection != "drone"]
    drone_rows = [r for r in all_rows if r.source_collection == "drone"]

    document = Document()

    title = document.add_heading(
        f"Vegetation Indices Report — {field_name}", level=0)
    for run in title.runs:
        run.font.size = Pt(18)

    meta = document.add_paragraph()
    run = meta.add_run(
        f"Generated {datetime.now().strftime('%d %b %Y at %H:%M')}")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    run.italic = True

    document.add_paragraph()

    # Satellite
    document.add_heading("Satellite", level=1)
    if satellite_rows:
        start = min(r.satellite_image_date for r in satellite_rows)
        end = max(r.satellite_image_date for r in satellite_rows)
        for key, label, color in INDEX_ORDER:
            points = _build_weekly_trend(satellite_rows, key, start, end)
            _add_index_section(document, key, label, color, points)
    else:
        document.add_paragraph("No satellite data available for this field.")

    # Drone
    if drone_rows:
        document.add_heading("Drone", level=1)
        for key, label, color in INDEX_ORDER:
            points = _build_drone_series(drone_rows, key)
            if not any(p["count"] for p in points):
                continue
            _add_index_section(document, key, label, color, points)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
